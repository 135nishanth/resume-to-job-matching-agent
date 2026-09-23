import io
import re
from typing import Tuple
from pypdf import PdfReader
import docx


class DocumentParsingError(Exception):
    pass


def clean_text(text: str) -> str:
    """Normalize whitespace, remove null characters and clean line breaks."""
    if not text:
        return ""
    # Remove null characters
    text = text.replace("\x00", "")
    # Replace non-breaking spaces
    text = text.replace("\u00a0", " ")
    # Replace carriage returns
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Collapse multiple empty lines
    text = re.sub(r"\n{3,}", "\n\n", text)
    # Strip leading and trailing whitespace
    return text.strip()


def parse_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF bytes."""
    try:
        reader = PdfReader(io.BytesIO(file_bytes))
        extracted_pages = []
        for i, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                extracted_pages.append(page_text)
        
        full_text = "\n\n".join(extracted_pages)
        if not full_text.strip():
            raise DocumentParsingError("PDF file appears to be empty or contains scanned images without text.")
        return clean_text(full_text)
    except Exception as e:
        if isinstance(e, DocumentParsingError):
            raise
        raise DocumentParsingError(f"Failed to parse PDF document: {str(e)}")


def parse_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX bytes including paragraphs and tables."""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = []
        for p in doc.paragraphs:
            if p.text.strip():
                paragraphs.append(p.text)
        
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n".join(paragraphs)
        if not full_text.strip():
            raise DocumentParsingError("DOCX file appears to be empty.")
        return clean_text(full_text)
    except Exception as e:
        if isinstance(e, DocumentParsingError):
            raise
        raise DocumentParsingError(f"Failed to parse DOCX document: {str(e)}")


def parse_txt(file_bytes: bytes) -> str:
    """Extract text from raw bytes attempting multiple encodings."""
    encodings = ["utf-8", "utf-8-sig", "latin-1", "cp1252", "utf-16"]
    for enc in encodings:
        try:
            decoded = file_bytes.decode(enc)
            return clean_text(decoded)
        except (UnicodeDecodeError, LookupError):
            continue
    raise DocumentParsingError("Unable to decode text file. Ensure it is saved with UTF-8 encoding.")


def parse_document(file_bytes: bytes, filename: str) -> str:
    """
    Main parser dispatch based on filename extension.
    Validates size and file content.
    """
    if not file_bytes or len(file_bytes) == 0:
        raise DocumentParsingError("Provided document is completely empty.")

    # 15MB limit check
    if len(file_bytes) > 15 * 1024 * 1024:
        raise DocumentParsingError("File size exceeds maximum allowed limit of 15MB.")

    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        return parse_pdf(file_bytes)
    elif lower_name.endswith(".docx"):
        return parse_docx(file_bytes)
    elif lower_name.endswith(".txt") or lower_name.endswith(".md"):
        return parse_txt(file_bytes)
    else:
        # Fallback: attempt text decoding
        try:
            return parse_txt(file_bytes)
        except Exception:
            raise DocumentParsingError(
                f"Unsupported file format '{filename}'. Please upload a PDF (.pdf), Word document (.docx), or plain text (.txt) file."
            )
